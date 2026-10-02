#!/usr/bin/env python
"""Ground truth for Agilent Seahorse XF `.asyr` files (`agilent-seahorse-asyr`) — standard library
only. No independent reader of `.asyr` exists and no vendor export was published with these files,
so this is a SECOND IMPLEMENTATION of our format notes (`independent: false`): the corpus test
reports a match as self-consistency, not as validation.

Usage:  python seahorse_oracle.py [--out DIR] --id ID FILE

Recorded: the assay name, instrument serial, software version; the plate size, well names,
groups and background wells; the injections of every group (port, reagent, concentration, unit,
volume); the number of plate readings and measurements; the corrected emission of every analyte at
up to 64 (reading, well) points; and one physical check of the well order: during the first three
measurements the O2 emission of background wells (no cells) changes far less than that of the
other wells (`background_ratio` = mean |change| of background wells / of the others).

With `--export PRISM` (Wave's export of the same assay to GraphPad Prism, `.pzfx`: Prism XML or
Prism's binary file) the oracle also records WAVE'S OWN RATES — per data table (`OCR Data`,
`ECAR Data`, `PER Data`) each Y column's title and replicate values per measurement, the
measurement times, and the group layout Wave writes into the project notes (selected groups and
their wells, background wells) — and becomes `independent: true`: the vendor software computed
those numbers from the same readings. The binary Prism layout (a column title, `10 00 ?? ?? 00 00
01 00`, the cell count, then 16-byte cells: u32 flag, u32 number format, f64 value; rows ×
replicate slots) was read from these files, not from any GraphPad document.

`--wave-xlsx XLSX --id ID` (no `.asyr`) records a Wave Excel export instead (sheets `Raw`, `Rate`,
`Assay Configuration`, `Calibration`, read with openpyxl): per reading and well the corrected O2
and pH emissions and Wave's O2 and pH levels, the calibration emissions, the rate constants and
Wave's OCR, ECAR and PER; `crates/openreadout-corpus-tests/tests/seahorse_rates.rs` runs the
reader's level and rate functions on the emissions and compares.

Writes `<DIR>/<ID>.json`; `crates/openreadout-corpus-tests/tests/seahorse_oracle/mod.rs` compares.
"""
from __future__ import annotations

import gzip
import json
import sys
import xml.etree.ElementTree as ET
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import oracle_json  # noqa: E402  (<id>.json, or <id>.json.gz over 1 MiB in corpus/oracle/heldout)

HERE = Path(__file__).resolve().parent
OUT = HERE.parent / "corpus" / "oracle" / "seahorse"


def strip(t: str) -> str:
    return t.split("}", 1)[-1]


def find(e, name):
    if e is None:
        return None
    for c in e:
        if strip(c.tag) == name:
            return c
    return None


def text(e, name):
    c = find(e, name)
    return (c.text or "").strip() if c is not None else None


def well_name(r: int, c: int) -> str:
    return "ABCDEFGHIJKLMNOP"[r] + str(c + 1)


PRISM = "{http://graphpad.com/prism/Prism.htm}"


def _group_layout(notes: str) -> dict:
    """Selected groups (name -> wells) and background wells from Wave's project notes."""
    import re
    groups, background = {}, []
    for m in re.finditer(r"Group Name: (.*?)\s*Group Wells: ([^\x00-\x08\n\r<]*)", notes, re.S):
        name = m.group(1).strip()
        if name.startswith("["):
            continue  # unselected group: not exported
        clean = []
        for w in m.group(2).split(","):
            w = w.strip().rstrip("-").strip()
            if w.startswith("["):
                continue  # unselected well
            mm = re.fullmatch(r"[\(\{]?([A-P]\d{1,2})[\)\}]?", w)
            if mm:
                clean.append(mm.group(1))
        if name == "Background":
            background = clean
        else:
            groups[name] = clean
    return {"selected_groups": [{"name": k, "wells": v} for k, v in groups.items()], "background": background}


def _prism_xml(path: Path) -> dict:
    root = ET.parse(path).getroot()
    lines = []
    for n in root.iter(PRISM + "Notes"):
        for f in n.iter(PRISM + "Font"):
            lines.append("".join(f.itertext()))
    out = {"kind": "prism-xml", "tables": {}}
    out.update(_group_layout("\n".join(lines)))
    for tab in root.iter(PRISM + "Table"):
        title = "".join(tab.find(PRISM + "Title").itertext()).strip()
        key = {"OCR Data": "OCR", "ECAR Data": "ECAR", "PER Data": "PER"}.get(title)
        if not key:
            continue
        x = tab.find(PRISM + "XColumn")
        if x is not None:
            out["times_min"] = [float(d.text) for d in x.iter(PRISM + "d")]
        cols = []
        for y in tab.findall(PRISM + "YColumn"):
            reps = [[float(d.text) if d.text else None for d in sc.findall(PRISM + "d")] for sc in y.findall(PRISM + "Subcolumn")]
            cols.append({"title": "".join(y.find(PRISM + "Title").itertext()).strip(), "replicates": reps})
        out["tables"][key] = cols
    return out


def _prism_binary(path: Path) -> dict:
    import re
    import struct
    data = path.read_bytes()
    if not data.startswith(b"PCFFGRA4"):
        raise SystemExit(f"{path}: not a Prism file")
    text_all = data.decode("latin-1")
    if "ECAR Data" in text_all or "PER Data" in text_all:
        raise SystemExit(f"{path}: more than the OCR table; this reading of binary Prism files only knows one-table exports")
    out = {"kind": "prism-binary", "tables": {"OCR": []}}
    out.update(_group_layout(text_all.replace("\r", "\n")))
    cols = []
    for m in re.finditer(rb"\x09\x00(.{4})", data, re.S):
        tl = struct.unpack("<I", m.group(1))[0]
        s = m.end()
        if not (1 <= tl <= 200) or s + tl + 12 > len(data):
            continue
        title = data[s:s + tl]
        if not title.endswith(b"\x00") or not all(32 <= b < 127 for b in title[:-1]):
            continue
        h = s + tl
        if data[h:h + 2] != b"\x10\x00" or data[h + 4:h + 8] != b"\x00\x00\x01\x00":
            continue
        n = struct.unpack_from("<I", data, h + 8)[0]
        c0 = h + 12
        if n == 0 or c0 + 16 * n > len(data):
            continue
        cells = [struct.unpack_from("<IId", data, c0 + 16 * k) for k in range(n)]
        cols.append((title[:-1].decode().strip(), cells))
    times = [c for c in cols if c[0].startswith("Time")]
    if not times:
        raise SystemExit(f"{path}: no time column")
    nrows = len(times[0][1])
    out["times_min"] = [v for _, _, v in times[0][1]]
    for title, cells in cols:
        if title.startswith("Time") or len(cells) % nrows:
            continue
        slots = len(cells) // nrows
        rows = [cells[r * slots:(r + 1) * slots] for r in range(nrows)]
        used = max((i + 1 for r in rows for i, c in enumerate(r) if c != (0, 0, 0.0)), default=0)
        if used == 0:
            continue
        reps = [[rows[r][i][2] for r in range(nrows)] for i in range(used)]
        out["tables"]["OCR"].append({"title": title, "replicates": reps})
    return out


def _wave_xlsx(path: Path, ident: str) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows = list(wb["Raw"].iter_rows(values_only=True))
    hdr = {h: i for i, h in enumerate(rows[0]) if h}
    data = [r for r in rows[1:] if r[0] is not None]
    wells = []
    for r in data:
        if r[hdr["Well"]] not in wells:
            wells.append(r[hdr["Well"]])
    nw = len(wells)
    nt = len(data) // nw

    def grid(name):
        return [[data[t * nw + w][hdr[name]] for w in range(nw)] for t in range(nt)]

    def secs(v):
        h, m, s = str(v).split(":")
        return int(h) * 3600 + int(m) * 60 + float(s)

    conf = {}
    conf_rows = [list(r) for r in wb["Assay Configuration"].iter_rows(values_only=True)]
    for r in conf_rows:
        vals = [c for c in r if c is not None]
        if len(vals) >= 2 and isinstance(vals[0], str):
            conf.setdefault(vals[0], vals[1:])
    buffer = {}
    for i, r in enumerate(conf_rows):
        if r and r[0] == "Buffer Factor :":
            cols = {j: c for j, c in enumerate(r) if isinstance(c, int)}
            for k in range(1, 9):
                row = conf_rows[i + k]
                letter = next((c for c in row[:2] if isinstance(c, str)), None)
                for j, col in cols.items():
                    if letter and j < len(row) and row[j] is not None:
                        buffer["%s%02d" % (letter, col)] = row[j]
    cal_rows = [list(r) for r in wb["Calibration"].iter_rows(values_only=True)]
    ph_cal = {}
    for i, r in enumerate(cal_rows):
        for j, c in enumerate(r):
            if c == "pH Emission":
                for k in range(1, 9):
                    rr = cal_rows[i + k]
                    for col in range(1, 13):
                        ph_cal["%s%02d" % (rr[j + 1], col)] = rr[j + 1 + col]
    rate = {}
    rr = list(wb["Rate"].iter_rows(values_only=True))
    rh = {h: i for i, h in enumerate(rr[0]) if h}
    for r in rr[1:]:
        if r[0] is None:
            continue
        rate.setdefault(r[rh["Well"]], {})[int(r[rh["Measurement"]])] = [r[rh["Time"]], r[rh["OCR"]], r[rh["ECAR"]], r[rh["PER"]]]
    nm = max(max(v) for v in rate.values())
    # Wave's background wells are the ones its Rate sheet gives no rates (all zero)
    background = [w for w in wells if all(v[1] == 0 and v[2] == 0 for v in rate[w].values())]
    return {
        "id": ident,
        "format": "agilent-seahorse-asyr",
        "independent": True,
        "reader": "openpyxl on a Wave Excel export (sheets Raw, Rate, Assay Configuration, Calibration)",
        "software": conf.get("XFe Assay Version", [None])[0],
        "wells": wells,
        "background": background,
        "measurement": [int(data[t * nw][hdr["Measurement"]]) for t in range(nt)],
        "time_s": [secs(data[t * nw][hdr["TimeStamp"]]) for t in range(nt)],
        "o2_emission": grid("O2 Corrected Em."),
        "o2_level": grid("O2 (mmHg)"),
        "ph_emission": grid("pH Corrected Em."),
        "ph_level": grid("pH"),
        "ph_calibration_emission": [ph_cal.get(w) for w in wells],
        "buffer_factor": [buffer.get(w) for w in wells],
        "constants": {
            "ksv": conf["ksv"][0], "fo": conf["Calculated FO"][0], "chamber_ul": conf["Pseudo Volume"][0],
            "tau_ac": conf["TAC"][0], "tau_aw": conf["TAW"][0], "tau_w": conf["TW"][0], "tau_c": conf["TC"][0], "tau_p": conf["TP"][0],
            "gain": conf["Gain Equation"], "ph_cal": conf["Calibration pH"][0], "plate_volume": conf["Plate Volume"][0], "kvol": conf["kVol"][0],
            "ecar_technique": conf["ECAR Technique"][0], "edge_offset": conf["Edge Offset"][0], "calculation_method": conf["Calculation Method"][0],
        },
        "rates": {w: [rate[w][m] for m in range(1, nm + 1)] for w in wells},
    }


def main(argv: list[str]) -> int:
    out_dir, ident, files, export, xlsx = OUT, None, [], None, None
    it = iter(argv)
    for a in it:
        if a == "--out":
            out_dir = Path(next(it))
        elif a == "--id":
            ident = next(it)
        elif a == "--export":
            export = Path(next(it))
        elif a == "--wave-xlsx":
            xlsx = Path(next(it))
        else:
            files.append(Path(a))
    if xlsx is not None and ident:
        out = _wave_xlsx(xlsx, ident)
        out_dir.mkdir(parents=True, exist_ok=True)
        (out_dir / f"{ident}.json").write_text(json.dumps(out, ensure_ascii=False) + "\n")
        print(ident, len(out["wells"]), "wells", len(out["time_s"]), "readings, background", out["background"])
        return 0
    if len(files) != 1 or not ident:
        print(__doc__, file=sys.stderr)
        return 2
    root = ET.fromstring(gzip.decompress(files[0].read_bytes()))
    plate = find(root, "Plate")
    wells, groups, background, injections = [], [], [], []
    seen = set()
    for w in find(plate, "Wells"):
        name = well_name(int(text(w, "RowIndex")), int(text(w, "ColumnIndex")))
        g = find(w, "ParentGroup")
        gname = text(g, "GroupName") if g is not None else None
        wells.append(name)
        groups.append(gname or "")
        if g is not None and text(g, "IsBackground") == "true":
            background.append(name)
        if gname and gname not in seen:
            seen.add(gname)
            inj = find(find(g, "InjectionCondition"), "Injections")
            for p in list(inj) if inj is not None else []:
                vals = {strip(x.tag): (x.text or "").strip() for x in p.iter() if (x.text or "").strip()}
                if not vals.get("ReagentName") and not vals.get("PortLocation"):
                    continue
                injections.append({"group": gname, "port": vals.get("PortLocation", ""), "reagent": vals.get("ReagentName", ""),
                                   "concentration": float(vals["PortConcentration"]) if vals.get("PortConcentration") else None,
                                   "unit": vals.get("PortConcentrationUnit", ""),
                                   "volume": float(vals["Volume"]) if vals.get("Volume") else None})
    ads = find(root, "AssayDataSet")
    ticks = list(find(ads, "PlateTickDataSets"))
    spans = [(int(text(s, "StartTickIndex")), int(text(s, "EndTickIndex"))) for s in find(ads, "RateSpans")]

    def corrected(t, analyte):
        for item in find(t, "AnalyteDataSetsByAnalyteName"):
            if find(find(item, "Key"), "string").text == analyte:
                ds = find(find(item, "Value"), "AnalyteDataSet")
                return [float(x.text) for x in find(ds, "CorrectedEmissionValues")]
        return None

    analytes = sorted({find(find(i, "Key"), "string").text for i in find(ticks[0], "AnalyteDataSetsByAnalyteName")})
    n_t, n_w = len(ticks), len(wells)
    pts = sorted({(round(k * (n_t - 1) / 7), round(j * (n_w - 1) / 7)) for k in range(8) for j in range(8)})
    samples = {an: [[t, w, corrected(ticks[t], an)[w]] for t, w in pts] for an in analytes}
    # physical check of the well order: background wells barely change O2 emission
    delta = [0.0] * n_w
    for s, e in spans[:3]:
        a, b = corrected(ticks[s], "O2"), corrected(ticks[e], "O2")
        for i in range(n_w):
            delta[i] += abs(b[i] - a[i]) / 3
    bg = [i for i, w in enumerate(wells) if w in background]
    others = [i for i in range(n_w) if i not in bg]
    ratio = (sum(delta[i] for i in bg) / len(bg)) / (sum(delta[i] for i in others) / len(others)) if bg and others else None
    out = {
        "id": ident,
        "format": "agilent-seahorse-asyr",
        "independent": False,
        "reader": "Python standard library (gzip, xml.etree): a second implementation of docs/formats/agilent-seahorse.md",
        "name": text(root, "Name"),
        "serial": text(root, "InstrumentSerialNumber"),
        "sw_version": text(root, "SWVersion"),
        "rows": int(text(plate, "RowCount")),
        "cols": int(text(plate, "ColumnCount")),
        "wells": wells,
        "groups": groups,
        "background": background,
        "injections": injections,
        "readings": n_t,
        "measurements": len(spans),
        "measure_commands": sum(1 for c in find(ads, "CommandHistory") if text(c, "CommandName") == "Measure"),
        "analytes": analytes,
        "samples": samples,
        "background_ratio": ratio,
    }
    if export is not None:
        head = export.read_bytes()[:8]
        wave = _prism_binary(export) if head.startswith(b"PCFFGRA4") else _prism_xml(export)
        wave["file"] = export.name
        out["independent"] = True
        out["reader"] = ("Python standard library: a second implementation of docs/formats/agilent-seahorse.md for the structure; "
                         "Wave's own OCR/ECAR/PER from its Prism export of the same assay (" + wave["kind"] + ")")
        out["wave_export"] = wave
    out_dir.mkdir(parents=True, exist_ok=True)
    oracle_json.write_text(out_dir / f"{ident}.json", json.dumps(out, indent=1, ensure_ascii=False) + "\n")
    print(ident, n_w, "wells", n_t, "readings", len(spans), "measurements", "background ratio", ratio)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
