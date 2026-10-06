#!/usr/bin/env python
"""Ground truth for readers of sampled-column files (EPR, X-ray diffraction, electrochemistry,
thermal analysis): `crates/openreadout-corpus-tests/tests/series_oracle/mod.rs` compares.

Sources, never our reader:
- the vendor software's export of the same measurement (`--export`): an `.xy`/`.txt`/`.csv`/`.asc`
  table parsed with the standard library (`--skip-until`, `--xcol`, `--ycol`, `--delim`,
  `--decimal-comma`, `--y-mult`);
- DeerLab's BES3T reader (MIT, vendored unchanged in `third_party/deerload.py`) on `.DSC/.DTA`
  (`--deerload`): every sweep's real (and imaginary) values;
- values the depositor states outside the file (`--fact path=value[:rel_tol]`);
- TRIOS exports of a `.tri` run (`--trios-export PATH [--trace K] [--sample-mass-g M] [--sheet S]`):
  the Excel export, the CSV export or a depositor's copy into a workbook;
- neware_reader (BSD-2-Clause) on an older Neware `.nda` that NewareNDA refuses (`--neware-reader`).

Usage:
  python series_oracle.py --id ID --format FMT FILE [--out DIR]
      [--export PATH --parse table [--skip-until TEXT] [--xcol 0] [--ycol 1] [--delim ,]
       [--decimal-comma] [--trace 0] [--sweep 0] [--channel intensity] [--x-tol 1e-6]
       [--y-tol-rel 1e-6] [--y-tol-abs 0] [--y-mult 1] [--first-row 0] [--x-mult 1]] ...
      [--deerload [--x-tol 1e-6]]
      [--fact method.parameters.scans=16] ...

Writes `<DIR>/<ID>.json` (default `corpus/oracle/series/`). Rows are sampled at up to 256 evenly
spaced indices.
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parent))
import oracle_json  # noqa: E402  (<id>.json, or <id>.json.gz over 1 MiB in corpus/oracle/heldout)

HERE = Path(__file__).resolve().parent
OUT = HERE.parent / "corpus" / "oracle" / "series"
ROWS = 256


def pick(n: int) -> list[int]:
    if n <= 0:
        return []
    return sorted({round(i * (n - 1) / (ROWS - 1)) for i in range(ROWS)}) if n > ROWS else list(range(n))


def num(s: str, comma: bool) -> float | None:
    s = s.strip().strip('"')
    if comma:
        s = s.replace(".", "").replace(",", ".") if s.count(",") == 1 and s.count(".") > 0 else s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def parse_table(path: Path, a: dict) -> tuple[list[float | None], list[float]]:
    text = path.read_bytes().decode("utf-8", "replace").lstrip("﻿")
    lines = text.splitlines()
    if a.get("skip_until"):
        for k, l in enumerate(lines):
            if a["skip_until"] in l:
                lines = lines[k + 1:]
                break
        else:
            raise SystemExit(f"{path.name}: marker {a['skip_until']!r} not found")
    comma = a.get("decimal_comma", False)
    delim = a.get("delim")
    xs, ys = [], []
    xc, yc = a.get("xcol"), int(a.get("ycol", 1))
    for l in lines:
        if not l.strip():
            continue
        f = l.split(delim) if delim else re.split(r"[\s,;]+" if not comma else r"[\s;]+", l.strip())
        f = [c for c in f if c.strip() != ""] if not delim else f
        try:
            y = num(f[yc], comma)
            x = num(f[int(xc)], comma) if xc is not None else None
        except IndexError:
            continue
        if y is None or (xc is not None and x is None):
            continue
        xs.append(x)
        ys.append(y)
    first = int(a.get("first_row", 0))
    xs, ys = xs[first:], ys[first:]
    return xs, ys


def parse_values(path: Path, a: dict) -> tuple[list[float | None], list[float]]:
    """Every number after the `--skip-until` marker up to `--stop-at` (PDXL `.asc` data blocks)."""
    text = path.read_bytes().decode("latin-1")
    lines = text.splitlines()
    marker = re.compile(re.escape(a["skip_until"]) + r"\s*=")
    start = next(k for k, l in enumerate(lines) if marker.match(l)) + 1
    ys = []
    for l in lines[start:]:
        if a.get("stop_at") and l.startswith(a["stop_at"]):
            break
        ys += [float(v) for v in re.split(r"[\s,]+", l.strip()) if v]
    return [None] * len(ys), ys


def export_trace(path: Path, a: dict) -> dict:
    xs, ys = parse_values(path, a) if a.get("parse") == "values" else parse_table(path, a)
    ym = float(a.get("y_mult", 1.0))
    xm = float(a.get("x_mult", 1.0))
    idx = pick(len(ys))
    return {
        "trace": int(a.get("trace", 0)),
        "sweep": int(a.get("sweep", 0)),
        "channel": a.get("channel", "intensity"),
        "n": len(ys) if not a.get("no_count") else None,
        "samples": [[i, None if xs[i] is None else xs[i] * xm, ys[i] * ym] for i in idx],
        "x_tol": float(a.get("x_tol", 1e-6)),
        "y_tol_rel": float(a.get("y_tol_rel", 1e-6)),
        "y_tol_abs": float(a.get("y_tol_abs", 0.0)),
        "source": f"export {path.name}",
    }


def deerload_traces(path: Path, a: dict) -> list[dict]:
    sys.path.insert(0, str(HERE / "third_party"))
    import numpy as np
    from deerload import deerload  # type: ignore

    ab, data, pars = deerload(str(path), full_output=True)
    desc = pars["DESC"]
    x = ab[0] if isinstance(ab, list) else ab
    # deerload's key/value pattern `(\w+)\W+(.*)` swallows a leading minus sign of the value, so
    # a negative XMIN comes back positive: its abscissa is not used then (reported upstream as a
    # black-box finding; the values are unaffected)
    x_ok = desc.get("XTYP") == "IDX" and not str(desc.get("XMIN", "")).strip().startswith("-")
    raw_dsc = Path(str(path)[:-4] + ".DSC").read_text(errors="replace")
    m = re.search(r"^XMIN\s+(\S+)", raw_dsc, re.M)
    if m and m.group(1).startswith("-"):
        x_ok = False
    x = np.asarray(x) * 1e3  # deerload divides every abscissa by 1000 (ns -> µs)
    data = np.asarray(data)
    if data.ndim == 1:
        data = data[:, None]
    cplx = np.iscomplexobj(data)
    out = []
    ny = data.shape[1]
    sweeps = sorted({0, ny - 1, ny // 2}) if ny > 1 else [0]
    for s in sweeps:
        col = data[:, s]
        idx = pick(len(col))
        for ch, vals in (("real", col.real), ("imaginary", col.imag)) if cplx else (("intensity", col),):
            out.append({
                "trace": 0,
                "sweep": int(s),
                "sweeps": int(ny),
                "channel": ch,
                "n": int(len(col)),
                "samples": [[i, float(x[i]) if x_ok else None, float(vals[i])] for i in idx],
                "x_tol": float(a.get("x_tol", 1e-6)),
                "y_tol_rel": 1e-12,
                "y_tol_abs": 0.0,
                "source": "DeerLab deerload (MIT, black box)",
            })
    return out


def mpt_traces(path: Path, a: dict) -> list[dict]:
    """Every numeric column of an EC-Lab `.mpt` export, addressed by its label (the comparator
    finds the channel whose `extra.label` is it; `optional` columns may be absent from a binary
    file, which does not store what EC-Lab computes at export). Harmonic-analysis columns skip
    the rows EC-Lab blanks (-1, 0)."""
    t = path.read_bytes().decode("latin-1")
    lines = t.splitlines()
    nh = 1
    if lines and lines[0].startswith("EC-Lab ASCII FILE"):
        nh = int(lines[1].split(":")[1])
    # EC-Lab writes an empty label between some columns; the rows leave it out
    raw = [h.strip() for h in lines[nh - 1].split("\t")]
    rows = [l.rstrip("\t\r").split("\t") for l in lines[nh:] if l.strip()]
    head = [h for h in raw if h]
    if "" in raw and rows and len(rows[0]) == raw.index(""):
        head = [h for h in raw[:raw.index("")] if h]
    comma = bool(rows) and "," in "\t".join(rows[0]) and "." not in "\t".join(rows[0])
    def f(v):
        return float(v.replace(",", ".") if comma else v)
    cols = {}
    for k, h in enumerate(head):
        try:
            cols[h] = [f(r[k]) for r in rows]
        except ValueError:  # e.g. time/s exported as absolute date-times ("07/10/2022 01:08:59.8225"): not compared
            continue
    time = cols.get("time/s")
    idx = pick(len(rows))
    out = []
    for h, vals in cols.items():
        harm = h.startswith(("THD", "NSD", "NSR", "|Ewe h", "|I h"))
        samples = [[i, time[i] if time else None, vals[i]] for i in idx if not (harm and vals[i] in (-1.0, 0.0))]
        out.append({"trace": 0, "sweep": 0, "label": h, "optional": bool(a.get("optional")), "n": len(rows),
                    "samples": samples, "x_tol": 1e-9, "y_tol_rel": float(a.get("y_tol_rel", 1e-7)),
                    "y_tol_abs": float(a.get("y_tol_abs", 1e-30)), "source": f"EC-Lab export {path.name}"})
    return out


def galvani_fields(mpr: Path) -> tuple[dict, str | None]:
    """galvani (GPL-3.0, run as a black box) on an `.mpr`: every column but the packed flags, and
    the acquisition start it reports."""
    from galvani import BioLogic  # type: ignore

    m = BioLogic.MPRfile(str(mpr))
    cols = {k: [float(v) for v in m.data[k]] for k in m.data.dtype.names if k != "flags"}
    ts = getattr(m, "timestamp", None)
    return cols, (ts.isoformat() if ts else None)


def galvani_traces(target: Path, a: dict) -> tuple[list[dict], list[dict]]:
    """Columns of the `.mpr` (`--galvani` on an `.mpr` input, or `--galvani PAIRED.mpr` for an
    `.mpt` input), compared by label; `required` labels (those EC-Lab's export also names) must
    be found, galvani-only spellings may be absent."""
    mpr = Path(a["file"]) if a.get("file") else target
    try:
        cols, ts = galvani_fields(mpr)
    except Exception as e:  # galvani does not know every column id: no second opinion then
        print(f"  galvani: {type(e).__name__}: {e}")
        return [], []
    required = set(a.get("required", []))
    t = cols.get("time/s")
    freq = cols.get("freq/Hz")
    out = []
    for h, vals in cols.items():
        idx = pick(len(vals))
        harm = h.startswith(("THD", "NSD", "NSR", "|Ewe h", "|I h"))
        # an .mpt input carries EC-Lab's export, which blanks harmonic analysis above ~100 kHz
        if harm and a.get("file") and freq:
            idx = [i for i in idx if freq[i] < 1e5]
        out.append({"trace": 0, "sweep": 0, "label": h, "optional": h not in required, "n": len(vals),
                    "samples": [[i, t[i] if t else None, vals[i]] for i in idx], "x_tol": 1e-9,
                    "y_tol_rel": 1e-7, "y_tol_abs": 1e-30, "source": f"galvani {mpr.name} (GPL, black box)"})
    facts = [{"path": "acquisition.started_at", "value": ts[:19], "prefix": True, "source": "galvani"}] if ts and a.get("start") else []
    return out, facts


def mpt_labels(path: Path) -> list[str]:
    lines = path.read_bytes().decode("latin-1").splitlines()
    nh = int(lines[1].split(":")[1]) if lines[0].startswith("EC-Lab ASCII FILE") else 1
    return [h.strip() for h in lines[nh - 1].split("\t") if h.strip()]


def gamry_traces(path: Path, a: dict) -> tuple[list[dict], list[dict]]:
    """gamry-parser (MIT, installed from PyPI, run as a black box): every curve it returns (all
    tables but OCVCURVE, in file order) and the OCV curve, compared by column label. The trace and
    sweep of each table follow the file's table order (CURVE, CURVE1… share one trace)."""
    import gamry_parser as gp  # type: ignore

    keys = []
    for line in path.read_bytes().decode("latin-1").splitlines():
        f = line.split("\t")
        if len(f) >= 2 and f[1].strip() == "TABLE":
            keys.append(f[0])
    groups: list[str] = []
    where = {}
    for k in keys:
        base = "CURVE" if re.fullmatch(r"CURVE\d*", k) else k
        if base not in groups:
            groups.append(base)
        where[k] = (groups.index(base), sum(1 for x in keys[:keys.index(k)] if ("CURVE" if re.fullmatch(r"CURVE\d*", x) else x) == base))
    p = gp.GamryParser(str(path))
    p.load()
    out = []
    curve_keys = [k for k in keys if k != "OCVCURVE"]
    n = p.get_curve_count()
    if n != len(curve_keys):
        raise ValueError(f"gamry-parser returns {n} curves, the file has {len(curve_keys)} tables")
    frames = [(curve_keys[i], p.get_curve_data(i)) for i in range(n)]
    if "OCVCURVE" in keys and p.ocv_exists:
        frames.append(("OCVCURVE", p.get_ocv_curve()))
    for key, df in frames:
        ti, sw = where[key]
        tcol = "T" if "T" in df.columns else ("Time" if "Time" in df.columns else None)
        for col in df.columns:
            try:
                vals = [float(v) for v in df[col]]
            except (TypeError, ValueError):
                continue
            idx = pick(len(vals))
            t = [float(v) for v in df[tcol]] if tcol else None
            out.append({"trace": ti, "sweep": sw, "label": col, "n": len(vals),
                        "samples": [[i, t[i] if t else None, vals[i]] for i in idx], "x_tol": 1e-9,
                        "y_tol_rel": 1e-12, "y_tol_abs": 0.0, "source": f"gamry-parser {path.name} (MIT, black box)"})
    return out, []


def geddes_traces(path: Path, a: dict) -> list[dict]:
    """geddes (MIT; installed from PyPI, run as a black box): scan `--scans` (comma list, default 0)."""
    import geddes  # type: ignore

    out = []
    for k in [int(v) for v in str(a.get("scans", "0")).split(",")]:
        p = geddes.read(str(path), index=k)
        x, y = list(p.x), list(p.y)
        idx = pick(len(y))
        out.append({
            "trace": k, "sweep": 0, "channel": "intensity", "n": len(y),
            "samples": [[i, x[i], y[i]] for i in idx],
            "x_tol": float(a.get("x_tol", 1e-6)), "y_tol_rel": 1e-9, "y_tol_abs": 0.0,
            "source": "geddes 1.0.0 (MIT, black box)",
        })
    return out


def newarenda_traces(path: Path, a: dict) -> tuple[list[dict], list[dict]]:
    """NewareNDA (BSD-3-Clause, installed from PyPI, run as a black box) on an `.nda`/`.ndax`:
    every column it returns, with the file's own cycle numbers (`software_cycle_number=False`).
    In split `.ndax` files (`.ndc` 11-17) the capacities and energies are compared on the logged
    records only (NewareNDA integrates current between them; we return NaN there): the logged
    indices come from NewareNDA's own `read_ndc` of `data_runInfo.ndc`. Auxiliary channels are
    compared file by file (NewareNDA's `read_ndc` of each `data_AUX_*.ndc`), by our channel name."""
    import tempfile
    import zipfile

    import NewareNDA  # type: ignore
    from NewareNDA.NewareNDAx import read_ndc  # type: ignore
    from NewareNDA.dicts import state_dict  # type: ignore

    df = NewareNDA.read(str(path), software_cycle_number=False, log_level="ERROR")
    n = len(df)
    codes = {v: k for k, v in state_dict.items()}
    logged = None
    aux: list[tuple[str, list[float]]] = []
    if str(path).lower().endswith(".ndax"):
        z = zipfile.ZipFile(path)
        with tempfile.TemporaryDirectory() as d:
            if "data_runInfo.ndc" in z.namelist():
                ri = read_ndc(z.extract("data_runInfo.ndc", d))
                logged = set(int(v) for v in ri["Index"])
            names = sorted(m for m in z.namelist() if m.startswith("data_AUX_") and m.endswith(".ndc"))
            counts: dict[str, int] = {}
            for m in names:
                parts = m[len("data_AUX_"):-4].split("_")
                ty = int(parts[1]) + 100
                base = {102: "aux_voltage", 103: "aux_temperature"}.get(ty, "aux")
                counts[base] = counts.get(base, 0) + 1
                ours = f"{base}_{counts[base]}" if base != "aux" else f"aux_{ty}_{counts[base]}"
                adf = read_ndc(z.extract(m, d))
                col = [c for c in adf.columns if c not in ("Index", "Aux")]
                if len(col) == 1:
                    aux.append((ours, [float(v) for v in adf[col[0]]]))
    index = [int(v) for v in df["Index"]]
    cols = {
        "voltage": "Voltage", "current": "Current(mA)", "step_time": "Time",
        "charge_capacity": "Charge_Capacity(mAh)", "discharge_capacity": "Discharge_Capacity(mAh)",
        "charge_energy": "Charge_Energy(mWh)", "discharge_energy": "Discharge_Energy(mWh)",
        "step_index": "Step_Index", "record": "Index",
    }
    if a.get("cycle", True) and not (df["Cycle"] == 0).all():
        cols["cycle"] = "Cycle"
    if a.get("time_total"):
        cols["time"] = cols.pop("step_time")
    if "T1" in df.columns and not str(path).lower().endswith(".ndax"):
        cols["temperature"] = "T1"
    out = []
    every = pick(n)
    logged_rows = [i for i in range(n) if logged is None or index[i] in logged]
    for ours, theirs in cols.items():
        vals = [float(v) for v in df[theirs]]
        rows = every
        if logged is not None and ours.endswith(("capacity", "energy")):
            rows = [logged_rows[k] for k in pick(len(logged_rows))]
        samples = [[i, None, vals[i]] for i in rows if vals[i] == vals[i]]
        out.append({"trace": 0, "sweep": 0, "channel": ours, "n": n, "samples": samples, "x_tol": 1e-9,
                    "y_tol_rel": 2e-6 if ours in ("voltage", "current", "step_time", "time") or ours.endswith(("capacity", "energy")) else 1e-9,
                    "y_tol_abs": 1e-9, "source": f"NewareNDA {path.name} (BSD-3-Clause, black box)"})
    status = [float(codes[str(s)]) for s in df["Status"]]
    out.append({"trace": 0, "sweep": 0, "channel": "step_type", "n": n, "samples": [[i, None, status[i]] for i in every],
                "x_tol": 1e-9, "y_tol_rel": 0.0, "y_tol_abs": 0.0, "source": "NewareNDA status names mapped back to their codes"})
    for ours, vals in aux:
        # read_ndc returns every slot of the last page; the file's records are the first n
        idx = pick(min(len(vals), n))
        out.append({"trace": 0, "sweep": 0, "channel": ours, "n": n, "samples": [[i, None, vals[i]] for i in idx if vals[i] == vals[i]],
                    "x_tol": 1e-9, "y_tol_rel": 1e-9, "y_tol_abs": 1e-9, "source": "NewareNDA read_ndc of the auxiliary file"})
    facts = []
    ts = df["Timestamp"].iloc[0] if n else None
    if ts is not None and a.get("start"):
        if getattr(ts, "tzinfo", None) is not None:
            ts = ts.tz_convert("UTC")
        facts.append({"path": "acquisition.started_at", "value": ts.isoformat()[:19], "prefix": True, "source": "NewareNDA first timestamp"})
    return out, facts


def ta_export_traces(path: Path, a: dict) -> tuple[list[dict], list[dict]]:
    """A TA Universal Analysis text export (UTF-16, `SigN` label lines, space-separated rows) of
    the same run: every exported signal, matched to our channel by its label; rows are matched by
    time (Universal Analysis may export every n-th point). Our time channel is in seconds, the
    export's in minutes."""
    raw = path.read_bytes()
    text = raw.decode("utf-16") if raw[:2] in (b"\xff\xfe", b"\xfe\xff") else raw.decode("latin-1")
    labels = {}
    rows = []
    head: dict[str, str] = {}
    for line in text.splitlines():
        f = line.split("\t")
        if len(f) > 1 and f[0] and f[0] not in head:
            head[f[0]] = "\t".join(f[1:]).strip()
        if re.fullmatch(r"Sig\d+", f[0]) and len(f) > 1:
            labels[int(f[0][3:])] = f[1].strip()
            continue
        parts = line.split()
        if labels and len(parts) == len(labels):
            try:
                rows.append([float(v) for v in parts])
            except ValueError:
                pass
    names = [labels[k] for k in sorted(labels)]
    tcol = next(k for k, n in enumerate(names) if n.startswith("Time"))
    idx = pick(len(rows))
    out = []
    for k, name in enumerate(names):
        is_time = k == tcol
        # `interp`: the export is resampled at its own times (modulated DSC): our values linearly
        # interpolated there, within 0.3 % of the column's largest magnitude
        if a.get("interp"):
            big = max(abs(r[k]) for r in rows) or 1.0
            out.append({"trace": 0, "sweep": 0, "label": name, "n": None, "interpolate": True,
                        "samples": [[None, rows[i][tcol] * 60.0, rows[i][k]] for i in idx],
                        "x_tol": 0.01, "y_tol_rel": 0.0, "y_tol_abs": 1e-5 if is_time else 3e-3 * big,
                        "y_scale": 1 / 60.0 if is_time else 1.0,
                        "source": f"Universal Analysis export {path.name} (resampled; interpolated)"})
            continue
        out.append({"trace": 0, "sweep": 0, "label": name, "n": None,
                    "samples": [[None, rows[i][tcol] * 60.0, rows[i][k]] for i in idx],
                    "x_tol": 0.01, "y_tol_rel": 1e-6, "y_tol_abs": 1e-6 if is_time else 1e-7,
                    "y_scale": 1 / 60.0 if is_time else 1.0,
                    "source": f"Universal Analysis export {path.name}"})
    facts = []
    src = f"Universal Analysis export {path.name} header"
    if head.get("Sample"):
        facts.append({"path": "sample.name", "value": head["Sample"], "source": src})
    if head.get("Operator"):
        facts.append({"path": "acquisition.operator", "value": head["Operator"], "source": src})
    size = head.get("Size", "").split()
    if len(size) == 2 and size[1] == "mg" and float(size[0]) > 0:
        facts.append({"path": "method.parameters.sample_mass", "value": float(size[0]), "source": src})
    if head.get("Date") and head.get("Time"):
        facts.append({"path": "acquisition.started_at", "value": f"{head['Date']}T{head['Time']}", "source": src})
    return out, facts


# TRIOS exports: column label -> (scale applied to OUR value to reach the export's unit, abs tol, rel tol)
TRIOS_COLUMNS = {
    "Time": (1 / 60.0, 0.0051, 0.0),                 # min, 2 decimals
    "Step time": (1.0, 0.0, 2e-5),
    "Temperature": (1.0, 0.0051, 0.0),              # °C, 2 decimals
    "Weight": (1e6, 0.00051, 0.0),                  # mg, 3 decimals; ours kg
    "Angular frequency": (1.0, 0.0, 2e-5),
    "Oscillation torque": (1e6, 0.0, 2e-5),         # µN·m; ours N·m
    "Raw phase": (180.0 / 3.141592653589793, 0.0, 2e-5),  # °; ours rad
    "Oscillation displacement": (1.0, 0.0, 2e-5),
    "Storage modulus": (1.0, 1e-6, 1e-3),           # computed: TRIOS's definitions
    "Loss modulus": (1.0, 1e-6, 1e-3),
    "Tan(delta)": (1.0, 1e-9, 1e-3),
    "Complex viscosity": (1.0, 1e-9, 1e-3),
}


def trios_export_traces(path: Path, a: dict) -> tuple[list[dict], list[dict]]:
    """A TRIOS export of the same run: the Excel export (`.xls`: a `Details` sheet, then one sheet
    per procedure step with data, in step order, rows = the step's points), the CSV export (`;`,
    decimal comma: one step, `--trace K`, rows = the step's first points; heat flow normalised by
    the sample size, `--sample-mass-g`, with the opposite sign) or a depositor's copy into a
    workbook (`.xlsx`, `--sheet`: blocks of Time/Temperature/Weight columns, one per step). Every
    column whose label TRIOS and our channel share is compared row by row (index = sample)."""
    ext = path.suffix.lower()
    out, facts = [], []
    src = f"TRIOS export {path.name}"

    def add(trace, label, rows):
        if label == "Heat Flow (Normalized)":
            m = float(a["sample_mass_g"])
            scale, abs_tol, rel = -1.0 / m, 0.00051, 0.0
            label = "Heat Flow"
        elif label in TRIOS_COLUMNS:
            scale, abs_tol, rel = TRIOS_COLUMNS[label]
        else:
            return
        idx = pick(len(rows))
        out.append({"trace": trace, "sweep": 0, "label": label, "n": None,
                    "samples": [[i, None, rows[i]] for i in idx if rows[i] is not None],
                    "y_scale": scale, "y_tol_abs": abs_tol, "y_tol_rel": rel,
                    "source": src})

    if ext == ".xls":
        import xlrd  # type: ignore
        wb = xlrd.open_workbook(str(path))
        sheets = wb.sheets()
        det = {sheets[0].cell_value(r, 0): sheets[0].cell_value(r, 1) for r in range(sheets[0].nrows)} if sheets[0].name == "Details" else {}
        for k, sh in enumerate(sheets[1:]):
            for c in range(sh.ncols):
                label = sh.cell_value(1, c)
                rows = []
                for r in range(3, sh.nrows):
                    v = sh.cell_value(r, c)
                    rows.append(float(v) if v != "" else None)
                add(k, label, rows)
        if det.get("Operator"):
            facts.append({"path": "acquisition.operator", "value": det["Operator"], "source": src + " Details sheet"})
        if det.get("Sample name"):
            facts.append({"path": "sample.name", "value": det["Sample name"], "source": src + " Details sheet"})
    elif ext == ".csv":
        lines = path.read_bytes().decode("latin-1").splitlines()
        head = [l.split(";") for l in lines[:3]]
        labels = head[1]
        rows = []
        for l in lines[3:]:
            f = l.split(";")
            try:
                rows.append([float(x.replace(",", ".")) for x in f[:len(labels)]])
            except ValueError:
                continue
        for c, label in enumerate(labels):
            add(int(a["trace"]), label.strip(), [r[c] for r in rows])
    elif ext == ".xlsx":
        import openpyxl  # type: ignore
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        rows = [list(r) for r in wb[a["sheet"]].iter_rows(values_only=True)]
        hdr_row = next(i for i, r in enumerate(rows) if r and "Time" in r and "Weight" in r)
        hdr = rows[hdr_row]
        starts = [j for j, v in enumerate(hdr) if v == "Time" and j + 2 < len(hdr) and hdr[j + 1] == "Temperature" and hdr[j + 2] == "Weight"]
        for k, j in enumerate(starts):
            for c, label in ((j, "Time"), (j + 1, "Temperature"), (j + 2, "Weight")):
                col = [r[c] if c < len(r) else None for r in rows[hdr_row + 2:]]
                while col and col[-1] is None:
                    col.pop()
                add(k, label, [float(v) if isinstance(v, (int, float)) else None for v in col])
    else:
        raise SystemExit(f"{path}: not a TRIOS export this oracle knows")
    return out, facts


def pyngb_traces(path: Path, a: dict) -> tuple[list[dict], list[dict]]:
    """pyNGB (MIT, installed from PyPI, run as a black box) on a NETZSCH `.ngb-*` file: every
    column of every run (the sample run is trace 0; the correction run of a sample + correction
    file is trace 1), compared by our channel name, and the metadata it reports. `--expdat XLSX`
    (repeatable) adds the depositor's Proteus export: at each exported time the sample
    temperature, within 0.6 K (Proteus exports calibrated, resampled temperatures)."""
    import json as _json

    import pyngb  # type: ignore

    rename = {f"environmental_acceleration_{c}": f"acceleration_{c}" for c in "xyz"}
    out: list[dict] = []
    facts: list[dict] = []
    runs = [("sample", 0)]
    first = pyngb.read_ngb(str(path))
    meta = _json.loads(first.schema.metadata[b"file_metadata"])
    try:
        pyngb.read_ngb(str(path), run="correction")
        if meta.get("measurement_type") in ("sample_correction",):
            runs.append(("correction", 1))
    except Exception:  # noqa: BLE001  (one run only)
        pass
    time0 = None
    for run, ti in runs:
        t = first if run == "sample" else pyngb.read_ngb(str(path), run=run)
        cols = {name: t.column(name).to_pylist() for name in t.column_names}
        time = cols.get("time")
        if ti == 0:
            time0 = time
        n = len(time) if time else 0
        idx = pick(n)
        for name, vals in cols.items():
            ours = rename.get(name, name)
            if re.fullmatch(r"[0-9a-f]{2}", ours):
                ours = f"channel_{ours}"
            out.append({"trace": ti, "sweep": 0, "channel": ours, "n": n,
                        "samples": [[i, time[i] if time else None, float(vals[i])] for i in idx if vals[i] is not None],
                        "x_tol": 1e-9, "y_tol_rel": 1e-12, "y_tol_abs": 0.0,
                        "source": f"pyNGB {path.name} (MIT, black box), run {run}"})
    for key, path_ in (("sample_mass", "method.parameters.sample_mass"), ("operator", "acquisition.operator"),
                       ("sample_name", "sample.name")):
        v = meta.get(key)
        # a correction run stores a sentinel mass (-1000 mg): no sample mass
        if v not in (None, "") and not (key == "sample_mass" and isinstance(v, (int, float)) and v <= 0):
            facts.append({"path": path_, "value": v, "source": "pyNGB metadata"})
    if meta.get("date_performed"):
        facts.append({"path": "acquisition.started_at", "value": str(meta["date_performed"])[:19], "prefix": True,
                      "source": "pyNGB date_performed"})
    for x in a.get("expdat", []):
        import openpyxl  # type: ignore

        wb = openpyxl.load_workbook(x, read_only=True)
        rows = [r for r in wb.worksheets[0].iter_rows(values_only=True)]
        head = [str(h) for h in rows[0]]
        ti_col = next(k for k, h in enumerate(head) if h.startswith("Time"))
        te_col = next(k for k, h in enumerate(head) if h.lstrip("#").startswith("Temp"))
        data = [r for r in rows[1:] if isinstance(r[ti_col], (int, float))]
        samples = []
        tsec = time0 or []
        for r in data[:: max(1, len(data) // 64)]:
            t = float(r[ti_col]) * 60.0
            i = min(range(len(tsec)), key=lambda k: abs(tsec[k] - t))
            samples.append([i, t, float(r[te_col])])
        out.append({"trace": 0, "sweep": 0, "channel": "sample_temperature", "n": None, "samples": samples,
                    "x_tol": 0.6, "y_tol_rel": 0.0, "y_tol_abs": 0.6,
                    "source": f"Proteus export {Path(x).name} (calibrated temperature at the exported times)"})
    return out, facts


def neware_reader_traces(path: Path) -> list[dict]:
    """neware_reader (BSD-2-Clause, github.com/FTHuld/neware_reader at a pinned commit, run as a black
    box) on an older `.nda` that NewareNDA refuses (version 8): voltage, current and step time per
    record, sampled at up to 256 records, compared by our channel name. It reports one capacity
    column for both directions, so capacities are not compared."""
    from neware_reader import neware  # type: ignore

    df = neware.read_nda(str(path))
    n = len(df)
    every = pick(n)
    out = []
    for ours, theirs in (("voltage", "voltage_V"), ("current", "current_mA"), ("step_time", "time_in_step")):
        vals = [float(v) for v in df[theirs]]
        out.append({"trace": 0, "sweep": 0, "channel": ours, "n": n,
                    "samples": [[i, None, vals[i]] for i in every if vals[i] == vals[i]],
                    "x_tol": 1e-9, "y_tol_rel": 2e-6, "y_tol_abs": 1e-9,
                    "source": f"neware_reader {path.name} (BSD-2-Clause, black box)"})
    return out


def main() -> None:
    args = sys.argv[1:]
    out_dir = OUT
    rid = fmt = None
    target = None
    exports: list[tuple[Path, dict]] = []
    facts = []
    deer = False
    ged: list[dict] = []
    mpts: list[dict] = []
    galv: list[dict] = []
    independent = True
    gam = False
    nda: list[dict] = []
    nwr = False
    ngbs: list[dict] = []
    tas: list[tuple[Path, dict]] = []
    trios: list[tuple[Path, dict]] = []
    cur: dict | None = None
    i = 0
    while i < len(args):
        k = args[i]
        if k == "--out":
            out_dir = Path(args[i + 1]); i += 2
        elif k == "--id":
            rid = args[i + 1]; i += 2
        elif k == "--format":
            fmt = args[i + 1]; i += 2
        elif k == "--export":
            cur = {}
            exports.append((Path(args[i + 1]), cur)); i += 2
        elif k == "--deerload":
            deer = True; cur = {}; deer_opts = cur; i += 1
        elif k == "--gamry":
            gam = True; i += 1
        elif k == "--ta-export":
            tas.append((Path(args[i + 1]), {})); i += 2
        elif k == "--ta-interp":
            tas[-1][1]["interp"] = True; i += 1
        elif k == "--trios-export":
            cur = {}
            trios.append((Path(args[i + 1]), cur)); i += 2
        elif k == "--pyngb":
            cur = {"expdat": []}; ngbs.append(cur); i += 1
        elif k == "--expdat":
            cur["expdat"].append(args[i + 1]); i += 2
        elif k == "--newarenda":
            cur = {}; nda.append(cur); i += 1
        elif k == "--neware-reader":
            nwr = True; i += 1
        elif k == "--no-cycle":
            cur["cycle"] = False; i += 1
        elif k == "--second-implementation":
            independent = False; i += 1
        elif k == "--galvani":
            cur = {}
            if i + 1 < len(args) and not args[i + 1].startswith("--") and args[i + 1].lower().endswith(".mpr"):
                cur["file"] = args[i + 1]; i += 1
            galv.append(cur); i += 1
        elif k == "--start":
            cur["start"] = True; i += 1
        elif k == "--mpt":
            cur = {"file": args[i + 1]}; mpts.append(cur); i += 2
        elif k == "--geddes":
            cur = {}; ged.append(cur); i += 1
        elif k == "--fact":
            p, v = args[i + 1].split("=", 1)
            tol = None
            if re.fullmatch(r"-?[\d.eE+-]+:[\d.eE+-]+", v):
                v, tol = v.split(":")
            try:
                val: object = float(v)
            except ValueError:
                val = v
            f = {"path": p, "value": val, "source": "depositor"}
            if tol:
                f["tol_rel"] = float(tol)
            facts.append(f); i += 2
        elif k in ("--decimal-comma", "--no-count", "--optional", "--time-total"):
            assert cur is not None
            cur[k[2:].replace("-", "_")] = True; i += 1
        elif k.startswith("--"):
            assert cur is not None, f"{k} before --export/--deerload"
            cur[k[2:].replace("-", "_")] = args[i + 1]; i += 2
        else:
            target = Path(args[i]); i += 1
    if not (rid and fmt and target):
        raise SystemExit(__doc__)
    try:
        traces = [export_trace(p, a) for p, a in exports]
        for t, ta_opts in tas:
            tt, tf = ta_export_traces(t, ta_opts)
            traces += tt
            facts += tf
        for t, t_opts in trios:
            tt, tf = trios_export_traces(t, t_opts)
            traces += tt
            facts += [f for f in tf if f not in facts]
        if deer:
            traces += deerload_traces(target, deer_opts)
        for g in ged:
            traces += geddes_traces(target, g)
        for m in mpts:
            traces += mpt_traces(Path(m["file"]), m)
        if gam:
            gt, _ = gamry_traces(target, {})
            traces += gt
        for o in ngbs:
            nt, nf = pyngb_traces(target, o)
            traces += nt
            facts += nf
        for o in nda:
            nt, nf = newarenda_traces(target, o)
            traces += nt
            facts += nf
        if nwr:
            traces += neware_reader_traces(target)
        for g in galv:
            # labels EC-Lab's export shares with galvani: required
            exp = next((Path(m["file"]) for m in mpts), target if str(target).lower().endswith((".mpt", ".txt")) else None)
            if exp:
                g["required"] = mpt_labels(exp)
            gt, gf = galvani_traces(target, g)
            stored = {x["label"] for x in gt}
            # an export column galvani also finds in the .mpr is stored: it must be read
            for x in traces:
                if x.get("label") and str(x.get("source", "")).startswith("EC-Lab export"):
                    x["optional"] = x["label"] not in stored
            traces += gt
            facts += gf
    except Exception as e:  # the independent reader could not read it: an oracle error, recorded
        out_dir.mkdir(parents=True, exist_ok=True)
        oracle_json.write_text(out_dir / f"{rid}.json", json.dumps({"id": rid, "format": fmt, "error": f"{type(e).__name__}: {e}"}, indent=1) + "\n")
        print(f"{rid}: oracle error")
        return
    o = {"id": rid, "format": fmt, "independent": independent, "traces": traces}
    if facts:
        o["facts"] = facts
    out_dir.mkdir(parents=True, exist_ok=True)
    oracle_json.write_text(out_dir / f"{rid}.json", json.dumps(o, indent=1) + "\n")
    print(f"{rid}: {len(traces)} trace comparisons, {len(facts)} facts")


if __name__ == "__main__":
    main()
