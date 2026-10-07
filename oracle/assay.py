#!/usr/bin/env python
"""Ground truth for the plate-reader assays of `openreadout analyze` (plate analysis), computed WITHOUT OpenReadout.

Usage:
    oracle/.venv/bin/python oracle/assay.py            # write corpus/oracle/assay/*.json
    oracle/.venv/bin/python oracle/assay.py --check    # exit 1 when a committed case drifts

Sources of truth, per case (each check names its source):

* SciPy (`scipy.optimize.curve_fit`, Levenberg–Marquardt; BSD-3) on the plate values read here
  with plain Python (`csv`, the Gen5 text) or openpyxl (MIT; SkanIt and Tecan workbooks).
* R `drc` 3.0-1 and `growthcurver` 0.3.1 (GPL; run as black boxes through `Rscript`).
  `OPENREADOUT_RSCRIPT` names the Rscript to use (default: `Rscript` on PATH); both packages must
  be installed there. Without R the drc/growthcurver checks are left out and `--check` fails.
* Vendor-computed results stored in the exports themselves: the SkanIt `Standard Curve` step
  (4PL parameters, back-calculated concentrations, `< Min`/`> Max` flags) and `Dilution Factor`
  step; the Gen5 `StdCurve`/`Curve Fitting Results` (linear and 4PL parameters) and
  `[Concentration]` results.
* Plain NumPy re-implementations of the documented kinetic definitions (book/src/guides/plate-analysis.md)
  for max slope, lag and AUC, where no third-party tool defines them the same way.

Each case is written to `corpus/oracle/assay/<case>.json`: the input (a corpus id or a fixture in
`crates/openreadout-assay/tests/fixtures`), the request (`AssayRequest` JSON), and checks
`{path, expect, rel, abs, source}` on the `AssayOutput` JSON. The Rust harness
(`crates/openreadout-assay/tests/oracle.rs`) runs them.
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import math
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import numpy as np
import scipy
from scipy.optimize import brentq, curve_fit
from scipy.stats import t as student_t

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "corpus" / "oracle" / "assay"
FIXTURES = ROOT / "crates" / "openreadout-assay" / "tests" / "fixtures"
CORPUS = Path(os.environ.get("OPENREADOUT_CORPUS_DIR", ROOT / "corpus" / "files"))
SCIPY = f"scipy {scipy.__version__} curve_fit"


# ------------------------------------------------------------------ helpers
def manifest_file(cid: str) -> Path:
    text = (ROOT / "corpus" / "manifest.toml").read_text(encoding="utf-8")
    m = re.search(r'\[\[file\]\]\nid = "' + re.escape(cid) + r'"\n(?:[^\[]*?\n)?filename = "([^"]+)"', text)
    if not m:
        raise SystemExit(f"{cid}: not in corpus/manifest.toml")
    return CORPUS / m.group(1)


def row_col(well: str) -> tuple[int, int]:
    m = re.fullmatch(r"([A-Z]+)0*(\d+)", well.strip().upper())
    r = 0
    for ch in m.group(1):
        r = r * 26 + ord(ch) - 64
    return r - 1, int(m.group(2)) - 1


def well_name(r: int, c: int) -> str:
    return f"{chr(65 + r)}{c + 1}"


def four_pl(x, a, b, c, d):
    x = np.asarray(x, dtype=float)
    return d + (a - d) / (1.0 + (x / c) ** b)


def five_pl(x, a, b, c, d, g):
    x = np.asarray(x, dtype=float)
    return d + (a - d) / (1.0 + (x / c) ** b) ** g


def fit(model, x, y, p0, sigma=None):
    x, y = np.asarray(x, float), np.asarray(y, float)
    p, cov = curve_fit(model, x, y, p0=p0, sigma=sigma, absolute_sigma=False, method="lm", maxfev=200000, xtol=1e-14, ftol=1e-14)
    w = np.ones_like(y) if sigma is None else 1.0 / np.asarray(sigma, float) ** 2
    r = y - model(x, *p)
    sse = float(np.sum(w * r * r))
    ybar = float(np.sum(w * y) / np.sum(w))
    sst = float(np.sum(w * (y - ybar) ** 2))
    return p, cov, sse, 1.0 - sse / sst


def inv4(y, a, b, c, d):
    t = (y - d) / (a - d)
    if not 0.0 < t < 1.0:
        return None
    return c * (1.0 / t - 1.0) ** (1.0 / b)


def inv5(y, a, b, c, d, g):
    t = (y - d) / (a - d)
    if not 0.0 < t < 1.0:
        return None
    u = t ** (-1.0 / g) - 1.0
    return c * u ** (1.0 / b) if u > 0 else None


def rscript() -> str | None:
    exe = os.environ.get("OPENREADOUT_RSCRIPT") or shutil.which("Rscript")
    if not exe:
        return None
    ok = subprocess.run([exe, "-e", 'suppressMessages({library(drc); library(growthcurver)})'], capture_output=True)
    return exe if ok.returncode == 0 else None


R = rscript()
R_VERSIONS = None


def run_r(code: str, data: dict) -> dict:
    with tempfile.TemporaryDirectory() as d:
        p = Path(d) / "in.json"
        p.write_text(json.dumps(data), encoding="utf-8")
        full = "suppressMessages({library(drc); library(growthcurver); library(jsonlite)})\n" f'd <- fromJSON("{p}")\n' + code
        res = subprocess.run([R, "-e", full], capture_output=True, text=True)
        if res.returncode != 0:
            raise SystemExit(f"R failed:\n{res.stderr[-3000:]}")
        out = res.stdout
        s = out.index("@@JSON@@") + 8
        return json.loads(out[s : out.index("@@END@@")])


def r_versions() -> str:
    global R_VERSIONS
    if R_VERSIONS is None:
        v = run_r('cat("@@JSON@@", toJSON(list(r=R.version.string, drc=as.character(packageVersion("drc")), gc=as.character(packageVersion("growthcurver"))), auto_unbox=TRUE), "@@END@@")', {})
        R_VERSIONS = v
    return R_VERSIONS


def drc_src() -> str:
    v = r_versions()
    return f"R drc {v['drc']} ({v['r']})"


def gc_src() -> str:
    v = r_versions()
    return f"R growthcurver {v['gc']} ({v['r']})"


DRC_FIT = r"""
df <- data.frame(x = d$x, y = d$y)
if (!is.null(d$w)) df$w <- d$w
args <- list(formula = y ~ x, data = df, fct = if (d$model == "LL.5") LL.5() else LL.4(),
             control = drmc(relTol = 1e-12, maxIt = 10000))
if (!is.null(d$w)) args$weights <- df$w
if (!is.null(d$start)) args$start <- d$start
m <- do.call(drm, args)
co <- coef(m); se <- sqrt(diag(vcov(m)))
names(co) <- sub(":.*", "", names(co)); names(se) <- names(co)
wt <- if (is.null(d$w)) 1 else d$w
out <- list(coef = as.list(co), se = as.list(se), rss = sum(((d$y - fitted(m)) * wt)^2))
if (d$model == "LL.4") {
  ed <- ED(m, 50, interval = "delta", display = FALSE)
  out$ed50 <- ed[1, 1]; out$ed50_se <- ed[1, 2]; out$ed50_lower <- ed[1, 3]; out$ed50_upper <- ed[1, 4]
}
if (!is.null(d$predict)) out$pred <- as.numeric(predict(m, newdata = data.frame(x = d$predict)))
cat("@@JSON@@", toJSON(out, digits = NA, auto_unbox = TRUE), "@@END@@")
"""


def drc(x, y, model="LL.4", w=None, predict=None, start=None) -> dict:
    """drc fit. drc's `weights` act as 1/σ (it minimises Σ (w r)²; found by black-box runs:
    with weights 1/|y| and a start on the same branch it reproduces SciPy's sigma = |y| fit).
    `start` in drc's order (b, c, d, e[, f])."""
    data = {"x": list(map(float, x)), "y": list(map(float, y)), "model": model}
    if w is not None:
        data["w"] = list(map(float, w))
    if predict is not None:
        data["predict"] = list(map(float, predict))
    if start is not None:
        data["start"] = list(map(float, start))
    return run_r(DRC_FIT, data)


NLS_FIT = r"""
df <- data.frame(x = d$x, y = d$y)
s <- d$start
m <- nls(y ~ dd + (aa - dd) / (1 + (x / cc)^bb), data = df,
         start = list(aa = s[1], bb = s[2], cc = s[3], dd = s[4]),
         control = nls.control(maxiter = 1000, tol = 1e-8, minFactor = 1e-12, warnOnly = TRUE))
out <- list(coef = as.list(coef(m)), se = as.list(sqrt(diag(vcov(m)))), rss = sum(residuals(m)^2))
cat("@@JSON@@", toJSON(out, digits = NA, auto_unbox = TRUE), "@@END@@")
"""


def nls4(x, y, start) -> dict:
    """Base R `nls` (Gauss–Newton), 4PL in our parameterisation, started at `start` (a, b, c, d)."""
    return run_r(NLS_FIT, {"x": list(map(float, x)), "y": list(map(float, y)), "start": list(map(float, start))})


def nls_src() -> str:
    return f"R stats::nls ({r_versions()['r']})"


GC_FIT = r"""
res <- list()
for (w in names(d$wells)) {
  g <- SummarizeGrowth(d$t, d$wells[[w]], bg_correct = d$bg)
  v <- g$vals
  res[[w]] <- list(k = v$k, n0 = v$n0, r = v$r, t_mid = v$t_mid, t_gen = v$t_gen, sigma = v$sigma, note = v$note)
}
cat("@@JSON@@", toJSON(res, digits = NA, auto_unbox = TRUE), "@@END@@")
"""


class Case:
    def __init__(self, name: str, inp: dict, request: dict, about: str):
        self.name, self.input, self.request, self.about = name, inp, request, about
        self.checks: list[dict] = []

    def add(self, path: str, expect, source: str, rel: float | None = None, abs_: float | None = None):
        c = {"path": path, "expect": expect if not isinstance(expect, (np.floating, np.integer)) else float(expect), "source": source}
        if rel is not None:
            c["rel"] = rel
        if abs_ is not None:
            c["abs"] = abs_
        self.checks.append(c)

    def doc(self) -> dict:
        return {"case": self.name, "about": self.about, "input": self.input, "request": self.request, "checks": self.checks}


def matrix(path: Path) -> dict[str, float]:
    """A plate-matrix CSV (`title,1,…,12` then `A,…`) as {well: value}."""
    rows = list(csv.reader(path.read_text(encoding="utf-8").splitlines()))
    cols = [int(c) for c in rows[0][1:]]
    out = {}
    for r in rows[1:]:
        if not r or not r[0]:
            continue
        for c, v in zip(cols, r[1:]):
            if v.strip():
                out[f"{r[0]}{c}"] = float(v)
    return out


# ------------------------------------------------------------------ synthetic: dose-response
def synth_dose_response() -> list[Case]:
    vals = matrix(FIXTURES / "assay-synth-dose-response.csv")
    lay = list(csv.DictReader((FIXTURES / "assay-synth-dose-response-layout.csv").read_text(encoding="utf-8").splitlines()))
    neg = np.array([vals[r["well"]] for r in lay if r["role"] == "negative"])
    pos = np.array([vals[r["well"]] for r in lay if r["role"] == "positive"])
    series: dict[str, list[tuple[float, float]]] = {}
    for r in lay:
        if r["role"] == "sample":
            series.setdefault(r["compound"], []).append((float(r["concentration"]), vals[r["well"]]))
    inp = {"fixture": "assay-synth-dose-response.csv", "corpus": "assay-synth-dose-response"}
    req = {"analysis": "dose-response", "layout": "fixture:assay-synth-dose-response-layout.csv"}
    raw = Case("synth-dose-response", inp, req, "4PL per compound on raw RLU; IC50, SE, CI, Hill, plateaus; Z′ of the controls")
    norm = Case("synth-dose-response-normalized", inp, dict(req, normalize="controls"), "4PL per compound on percent effect (controls): IC50, absolute IC50")
    # Z′ and S/B from the controls (sample SD)
    zp = 1 - 3 * (pos.std(ddof=1) + neg.std(ddof=1)) / abs(pos.mean() - neg.mean())
    raw.add("quality.z_prime", zp, "numpy (Zhang 1999 formula)", rel=1e-9)
    raw.add("quality.signal_to_background", neg.mean() / pos.mean(), "numpy", rel=1e-9)
    for name, pts in series.items():
        x = np.array([p[0] for p in pts])
        for case, y in ((raw, np.array([p[1] for p in pts])), (norm, 100 * (np.array([p[1] for p in pts]) - neg.mean()) / (pos.mean() - neg.mean()))):
            p0 = [y[x == x.min()].mean(), 1.0, float(np.exp(np.mean(np.log(x)))), y[x == x.max()].mean()]
            p, cov, sse, r2 = fit(four_pl, x, y, p0)
            a, b, c, d = p
            if b < 0:  # canonical b > 0
                a, d, b = d, a, -b
            se = np.sqrt(np.diag(cov))
            df = len(x) - 4
            tq = student_t.ppf(0.975, df)
            sel = f"compounds[compound={name}]"
            case.add(f"{sel}.ec50", c, SCIPY, rel=1e-5)
            case.add(f"{sel}.ec50_se", se[2], SCIPY + " (pcov)", rel=1e-4)
            lo, hi = math.exp(math.log(c) - tq * se[2] / c), math.exp(math.log(c) + tq * se[2] / c)
            case.add(f"{sel}.ec50_ci_low", lo, SCIPY + " + scipy.stats.t (log-scale Wald)", rel=1e-4)
            case.add(f"{sel}.ec50_ci_high", hi, SCIPY + " + scipy.stats.t (log-scale Wald)", rel=1e-4)
            case.add(f"{sel}.top", max(a, d), SCIPY, rel=1e-5)
            case.add(f"{sel}.bottom", min(a, d), SCIPY, rel=1e-5, abs_=1e-3)
            hill = b if d > a else -b
            case.add(f"{sel}.hill_slope", hill, SCIPY, rel=1e-5)
            case.add(f"{sel}.fit.r_squared", r2, SCIPY, rel=1e-8)
            if case is norm:
                case.add(f"{sel}.absolute_ec50", brentq(lambda v: four_pl(v, a, b, c, d) - 50.0, 1e-9, 1e6) if min(a, d) < 50 < max(a, d) else None, "scipy brentq on the curve_fit curve", rel=1e-5)
            if R:
                m = drc(x, y)
                # drc LL.4: f = c + (d − c)/(1 + exp(b (log x − log e))): e = EC50
                case.add(f"{sel}.ec50", m["ed50"], drc_src() + " ED(…, 50)", rel=1e-4)
                # drc's standard errors come from the Hessian of the RSS (observed information), ours
                # and SciPy's/nls's from the Gauss–Newton JᵀJ: they differ by the residual-curvature
                # term, a few percent when the fit is loose (CPD-C: 9 %)
                case.add(f"{sel}.ec50_se", m["ed50_se"], drc_src() + " ED(…, 50, interval = \"delta\"), Hessian-based SE", rel=0.1)
                case.add(f"{sel}.hill_slope", abs(m["coef"]["b"]) * (1 if hill > 0 else -1), drc_src() + " coef b", rel=1e-4)
                case.add(f"{sel}.top", max(m["coef"]["c"], m["coef"]["d"]), drc_src(), rel=1e-4)
                if case is raw:
                    n = nls4(x, y, [a, b, c, d])
                    case.add(f"{sel}.ec50", n["coef"]["cc"], nls_src(), rel=1e-6)
                    case.add(f"{sel}.ec50_se", n["se"]["cc"], nls_src() + " vcov (Gauss–Newton)", rel=1e-4)
    return [raw, norm]


# ------------------------------------------------------------------ synthetic: 5PL ELISA
def grid_layout(path: Path) -> dict[str, dict[str, str]]:
    out: dict[str, dict[str, str]] = {}
    title = None
    cols: list[int] = []
    for row in csv.reader(path.read_text(encoding="utf-8").splitlines()):
        if not row or not any(row):
            title = None
            continue
        if title is None:
            title, cols = row[0], [int(c) for c in row[1:]]
            out[title] = {}
            continue
        for c, v in zip(cols, row[1:]):
            if v:
                out[title][f"{row[0]}{c}"] = v
    return out


def synth_elisa() -> list[Case]:
    vals = matrix(FIXTURES / "assay-synth-elisa-5pl.csv")
    lay = grid_layout(FIXTURES / "assay-synth-elisa-5pl-layout.csv")
    role, sample = lay["role"], lay["sample"]
    conc = {w: float(v) for w, v in lay["concentration (pg/mL)"].items()}
    dil = {w: float(v.split(":")[1]) for w, v in lay["dilution"].items()}
    blank = np.mean([vals[w] for w, r in role.items() if r == "blank"])
    std = sorted((w for w, r in role.items() if r == "standard"), key=row_col)
    x = np.array([conc[w] for w in std])
    y = np.array([vals[w] - blank for w in std])
    inp = {"fixture": "assay-synth-elisa-5pl.csv", "corpus": "assay-synth-elisa-5pl"}
    req = {"analysis": "curve", "layout": "fixture:assay-synth-elisa-5pl-layout.csv", "model": "5pl", "weighting": "1/y2"}
    case = Case("synth-elisa-5pl", inp, req, "5PL standard curve weighted 1/y², back-calculated unknowns with dilution")
    p, cov, sse, r2 = fit(five_pl, x, y, [0.0, 1.4, 180.0, 3.2, 0.6], sigma=np.abs(y))
    se = np.sqrt(np.diag(cov))
    for i, n in enumerate("abcdg"):
        case.add(f"curve.fit.parameters[name={n}].value", p[i], SCIPY + " (sigma = |y|)", rel=1e-4, abs_=1e-6)
        case.add(f"curve.fit.parameters[name={n}].se", se[i], SCIPY + " (pcov)", rel=2e-3)
    case.add("curve.fit.r_squared", r2, SCIPY + " (weighted)", rel=1e-7)
    case.add("blank.value", blank, "numpy mean of H1, H2", rel=1e-12)
    groups: dict[str, list[str]] = {}
    for w, r in role.items():
        if r == "sample":
            groups.setdefault(sample[w], []).append(w)
    for name, ws in sorted(groups.items()):
        bcs = [inv5(vals[w] - blank, *p) for w in ws]
        if all(b is not None for b in bcs):
            for w, b in zip(ws, bcs):
                case.add(f"wells[well={w}].back_calculated", b, SCIPY + " inverse 5PL", rel=2e-4)
            d = dil.get(ws[0], 1.0)
            case.add(f"samples[group={name}].final_concentration", float(np.mean(bcs)) * d, SCIPY + " inverse 5PL × dilution", rel=2e-4)
    if R:
        # drc LL.5: weights are 1/σ (drc minimises Σ (w r)²), so 1/y² weighting is weights 1/|y|.
        # The 5PL has two asymmetric branches (sign of b); drc is started on SciPy's branch, in its
        # own parameter order (b, c = response at ∞, d = response at 0, e, f).
        a, b, c, d, g = p
        m = drc(x, y, "LL.5", w=1.0 / np.abs(y), predict=[conc[w] for w in std], start=[b, d, a, c, g])
        co = m["coef"]
        src_d = drc_src() + " LL.5, weights 1/|y|"
        for n, v in (("a", co["d"]), ("b", co["b"]), ("c", co["e"]), ("d", co["c"]), ("g", co["f"])):
            case.add(f"curve.fit.parameters[name={n}].value", v, src_d, rel=1e-4, abs_=1e-6)
        for w, dr in zip(std, m["pred"]):
            case.add(f"curve.fit.points[label={w}].fitted", dr, src_d, rel=1e-5)
        case.add("curve.fit.sse", m["rss"], src_d + ": Σ (w r)²", rel=1e-5)
    return [case]


# ------------------------------------------------------------------ synthetic: kinetics
def windows(t, y, h):
    best = None
    for s in range(0, len(t) - h + 1):
        slope, icpt = np.polyfit(t[s : s + h], y[s : s + h], 1)
        if best is None or slope > best[1]:
            best = (s, slope)
    return best


def kinetic_ref(t, y):
    """book/src/guides/plate-analysis.md definitions, re-implemented with NumPy."""
    n = len(t)
    h = max(5, math.ceil(n / 10))
    s, slope = windows(t, y, h)
    tc, yc = t[s : s + h].mean(), y[s : s + h].mean()
    base = y[: s + h].min()
    lag = tc - (yc - base) / slope if slope > 0 else None
    return {
        "max_slope_per_min": slope * 60,
        "lag_time_s": lag,
        "mean_slope_per_min": np.polyfit(t, y, 1)[0] * 60,
        "auc": float(np.trapezoid(y, t)),
        "time_to_max_s": float(t[int(np.argmax(y))]),
        "window": h,
    }


def synth_kinetics() -> list[Case]:
    rows = list(csv.DictReader((FIXTURES / "assay-synth-kinetics.csv").read_text(encoding="utf-8").splitlines()))
    wells: dict[str, list[tuple[float, float]]] = {}
    for r in rows:
        wells.setdefault(r["well"], []).append((float(r["time_min"]) * 60, float(r["value"])))
    case = Case("synth-kinetics", {"fixture": "assay-synth-kinetics.csv"}, {"analysis": "kinetics"},
                "max slope (Vmax), lag, mean slope, AUC, time to max per well")
    src = "numpy re-implementation of the documented definitions"
    for w, pts in wells.items():
        t = np.array([p[0] for p in pts])
        y = np.array([p[1] for p in pts])
        k = kinetic_ref(t, y)
        sel = f"kinetics[well={w}]"
        for f in ("max_slope_per_min", "lag_time_s", "mean_slope_per_min", "auc", "time_to_max_s", "window"):
            case.add(f"{sel}.{f}", k[f], src, rel=1e-9, abs_=1e-9)
    return [case]


# ------------------------------------------------------------------ growth (synthetic and Tecan)
def growth_ref(t_s, y):
    """µmax sliding window on ln(value) above 5 % of the max (book/src/guides/plate-analysis.md)."""
    n = len(t_s)
    h = max(5, math.ceil(n / 10))
    thr = 0.05 * max(y.max(), 0.0)
    keep = (y > thr) & (y > 0)
    lt, ly = t_s[keep] / 3600.0, np.log(y[keep])
    if len(lt) < h:
        return None
    s, mu = windows(lt, ly, h)
    return {"growth_rate_per_h": mu, "doubling_time_h": math.log(2) / mu}


def logistic(t, k, n0, r):
    return k / (1.0 + ((k - n0) / n0) * np.exp(-r * t))


def growth_checks(case: Case, t_s: np.ndarray, wells: dict[str, np.ndarray], bg: str):
    """`bg`: `none` (values already blank-subtracted) or `min` (subtract each well's minimum)."""
    t_h = (t_s - t_s[0]) / 3600.0
    src = "numpy re-implementation of the documented µmax definition"
    if R:
        gc = run_r(GC_FIT, {"t": list(map(float, t_h)), "wells": {w: list(map(float, v)) for w, v in wells.items()}, "bg": bg})
    for w, y in wells.items():
        yy = y - y.min() if bg == "min" else y
        sel = f"growth[well={w}]"
        ref = growth_ref(t_s, yy)
        if ref:
            case.add(f"{sel}.growth_rate_per_h", ref["growth_rate_per_h"], src, rel=1e-9)
            case.add(f"{sel}.doubling_time_h", ref["doubling_time_h"], src, rel=1e-9)
        # scipy logistic on the same corrected values
        try:
            p0 = [yy.max(), max(yy[0], yy.max() * 1e-3), 4.0 / t_h[-1]]
            p, cov = curve_fit(logistic, t_h, yy, p0=p0, bounds=([1e-12, 1e-12, 1e-12], [np.inf, np.inf, np.inf]), method="trf", x_scale="jac", xtol=1e-15, ftol=1e-15, gtol=1e-15, max_nfev=200000)
            case.add(f"{sel}.logistic_r_per_h", p[2], SCIPY + " logistic", rel=1e-4)
            case.add(f"{sel}.logistic_k", p[0], SCIPY + " logistic", rel=1e-5)
        except RuntimeError:
            pass
        if R and w in gc and gc[w]["note"] in ("", None):
            case.add(f"{sel}.logistic_r_per_h", gc[w]["r"], gc_src() + f" SummarizeGrowth(bg_correct = \"{bg}\")", rel=2e-3)
            case.add(f"{sel}.logistic_k", gc[w]["k"], gc_src(), rel=1e-3)
            case.add(f"{sel}.logistic_doubling_time_h", gc[w]["t_gen"], gc_src() + " t_gen", rel=2e-3)


def synth_growth() -> list[Case]:
    rows = list(csv.DictReader((FIXTURES / "assay-synth-growth.csv").read_text(encoding="utf-8").splitlines()))
    by: dict[str, list[tuple[float, float]]] = {}
    roles = {}
    for r in rows:
        by.setdefault(r["well"], []).append((float(r["time_h"]) * 3600, float(r["value"])))
        roles[r["well"]] = r["role"]
    t = np.array([p[0] for p in next(iter(by.values()))])
    blank = np.mean([[p[1] for p in by[w]] for w, ro in roles.items() if ro == "blank"], axis=0)
    wells = {w: np.array([p[1] for p in v]) - blank for w, v in by.items() if roles[w] != "blank"}
    case = Case("synth-growth", {"fixture": "assay-synth-growth.csv"}, {"analysis": "growth"},
                "growth rate and doubling time (ln-scale window) and logistic fit after per-time-point blank subtraction")
    growth_checks(case, t, wells, "none")
    return [case]


def tecan_growth() -> list[Case]:
    import openpyxl

    wb = openpyxl.load_workbook(manifest_file("tecan-icontrol-kinetic-xlsx"), data_only=True, read_only=True)
    rows = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
    i = next(k for k, r in enumerate(rows) if r and r[0] == "Time [s]")
    t = np.array([float(v) for v in rows[i][1:] if v is not None])
    wells = {}
    for r in rows[i + 2 :]:
        if not r or not isinstance(r[0], str) or not re.fullmatch(r"[A-H]\d{1,2}", r[0]):
            break
        v = np.array([float(x) for x in r[1 : 1 + len(t)]])
        if r[0] in ("A2", "B6", "C10", "D1", "E5", "F9", "G12", "H3"):
            wells[r[0]] = v
    case = Case("tecan-growth", {"corpus": "tecan-icontrol-kinetic-xlsx"}, {"analysis": "growth", "wells": "A2,B6,C10,D1,E5,F9,G12,H3"},
                "Tecan i-control kinetic OD600 (632 reads): per-well-minimum background, ln-scale µmax, logistic fit vs growthcurver")
    growth_checks(case, t, wells, "min")
    return [case]


# ------------------------------------------------------------------ SkanIt ELISA (vendor-computed)
def skanit_workbook():
    import openpyxl

    wb = openpyxl.load_workbook(manifest_file("skanit-elisa-steps"), data_only=True, read_only=True)

    def sheet_rows(prefix):
        ws = next(s for s in wb.worksheets if s.title.startswith(prefix))
        return [list(r) for r in ws.iter_rows(values_only=True)]

    def grid_after(rows, corner):
        i = next(k for k, r in enumerate(rows) if r and r[0] == corner)
        out = {}
        for r in rows[i + 1 : i + 9]:
            for c in range(12):
                v = r[c + 1]
                if v is not None and str(v).strip() != "":
                    out[f"{r[0]}{c + 1}"] = v
        return out

    return wb, sheet_rows, grid_after


def skanit_data():
    """The SkanIt ELISA read with openpyxl: raw absorbances, sample names, standard
    concentrations (µg/mL) and dilution factors from the `Layout definitions` sheet."""
    wb, sheet_rows, grid_after = skanit_workbook()
    raw = {w: float(v) for w, v in grid_after(sheet_rows("Absorbance 1"), "Abs").items()}
    names = grid_after(sheet_rows("Absorbance 1"), "Sample")
    # Layout definitions: row letter + names, then group line, then concentration / dilution line
    ld = [list(r) for r in next(s for s in wb.worksheets if s.title == "Layout definitions").iter_rows(values_only=True)]
    concs, dils = {}, {}
    for k, r in enumerate(ld):
        if r and isinstance(r[0], str) and re.fullmatch(r"[A-H]", r[0]):
            third = ld[k + 2]
            for c in range(12):
                v = third[c + 1]
                if isinstance(v, str) and "microg/ml" in v:
                    concs[f"{r[0]}{c + 1}"] = float(v.split()[0])
                elif isinstance(v, str) and re.fullmatch(r"1:\d+", v.strip()):
                    dils[f"{r[0]}{c + 1}"] = float(v.split(":")[1])
    return raw, names, concs, dils


def skanit_fit():
    """Blank-subtracted (mean of the Blank wells) unweighted 4PL on the standard wells."""
    raw, names, concs, dils = skanit_data()
    blank = float(np.mean([raw[w] for w, n in names.items() if str(n).startswith("Blank")]))
    std = sorted(concs, key=row_col)
    x = np.array([concs[w] for w in std])
    y = np.array([raw[w] - blank for w in std])
    p, cov, sse, r2 = fit(four_pl, x, y, [0.0, 1.0, 0.5, 3.0])
    return {"raw": raw, "names": names, "concs": concs, "dils": dils, "blank": blank, "x": x, "y": y, "p": p, "cov": cov, "r2": r2}


def skanit() -> list[Case]:
    wb, sheet_rows, grid_after = skanit_workbook()
    s = skanit_fit()
    raw, names, x, y, p, cov = s["raw"], s["names"], s["x"], s["y"], s["p"], s["cov"]
    case = Case("skanit-elisa", {"corpus": "skanit-elisa-steps"}, {"analysis": "curve"},
                "SkanIt ELISA with its own layout: blank subtraction, 4PL, back-calculated and dilution-corrected concentrations vs SkanIt's results")
    se = np.sqrt(np.diag(cov))
    for i, n in enumerate("abcd"):
        case.add(f"curve.fit.parameters[name={n}].value", p[i], SCIPY, rel=1e-5, abs_=1e-7)
        case.add(f"curve.fit.parameters[name={n}].se", se[i], SCIPY + " (pcov)", rel=1e-3)
    # vendor: the Standard Curve step's parameters
    sc = sheet_rows("Standard Curve 1")
    vend = {}
    for r in sc:
        if r and isinstance(r[0], str) and re.fullmatch(r"[abcd] =", r[0].strip()):
            vend[r[0].strip()[0]] = float(r[1])
    src_v = "SkanIt 7.0 Standard Curve step (vendor-computed, in the file)"
    for n in "abcd":
        case.add(f"curve.fit.parameters[name={n}].value", vend[n], src_v, rel=1e-3, abs_=1e-4)
    # vendor per-well concentrations and flags
    conc_grid = grid_after(sc, "Conc. [microg/ml]")
    for w, v in sorted(conc_grid.items(), key=lambda kv: row_col(kv[0])):
        if isinstance(v, (int, float)):
            case.add(f"wells[well={w}].back_calculated", float(v), src_v, rel=1.5e-3, abs_=2e-5)
        elif "Min" in str(v):
            case.add(f"wells[well={w}].flag", ["below_range", "below_curve"], src_v + " (< Min)")
        elif "Max" in str(v):
            case.add(f"wells[well={w}].flag", "above_range", src_v + " (> Max)")
    # vendor: Dilution Factor step, `Value` = concentration × dilution
    dfv = grid_after(sheet_rows("Dilution Factor 1"), "Value")
    src_d = "SkanIt 7.0 Dilution Factor step (vendor-computed)"
    for w, v in sorted(dfv.items(), key=lambda kv: row_col(kv[0])):
        if isinstance(v, (int, float)) and not math.isnan(float(v)) and not str(names.get(w, "")).startswith("Std"):
            case.add(f"wells[well={w}].final_concentration", float(v), src_d, rel=1.5e-3, abs_=1e-4)
    if R:
        m = drc(x, y)
        # drc LL.4: c ↔ our d (response at ∞ when b > 0), d ↔ our a, e ↔ our c
        co = m["coef"]
        ours = {"a": co["d"], "b": co["b"], "c": co["e"], "d": co["c"]} if co["b"] > 0 else {"a": co["c"], "b": -co["b"], "c": co["e"], "d": co["d"]}
        # this curve's top asymptote is barely reached (SE of c ≈ its value): drc's optimiser stops
        # in the flat valley with a slightly higher RSS than SciPy, nls, SkanIt and ours
        for n in "abcd":
            case.add(f"curve.fit.parameters[name={n}].value", ours[n], drc_src() + " LL.4 (stops early in a flat valley)", rel=5e-3, abs_=1e-5)
        nl = nls4(x, y, p)
        for i, n in enumerate("abcd"):
            case.add(f"curve.fit.parameters[name={n}].value", nl["coef"][{"a": "aa", "b": "bb", "c": "cc", "d": "dd"}[n]], nls_src(), rel=1e-6, abs_=1e-8)
            case.add(f"curve.fit.parameters[name={n}].se", nl["se"][{"a": "aa", "b": "bb", "c": "cc", "d": "dd"}[n]], nls_src() + " vcov (Gauss–Newton)", rel=1e-4)
    return [case]


# ------------------------------------------------------------------ Gen5 (vendor-computed)
def gen5_text(cid: str):
    """Plain-Python reading of a Gen5 text export: Layout (Well ID, Conc/Dil) and Results
    matrices (one line per read label), and the curve-fit table."""
    text = manifest_file(cid).read_bytes().decode("latin-1")
    lines = text.splitlines()
    layout: dict[str, dict[str, str]] = {}
    results: dict[str, dict[str, str]] = {}
    fit_rows = []
    section = None
    row = None
    for k, ln in enumerate(lines):
        cells = ln.split("\t")
        if ln.strip() in ("Layout", "Results"):
            section = ln.strip()
            row = None
            continue
        if ln.startswith("Curve Name"):
            fit_rows.append((cells, lines[k + 1].split("\t")))
            section = None
            continue
        if section and len(cells) >= 13 and cells[-1].strip():
            if cells[0] and re.fullmatch(r"[A-H]", cells[0]):
                row = cells[0]
            elif cells[0]:
                row = None
            if row:
                target = layout if section == "Layout" else results
                label = cells[-1].strip()
                for c in range(1, len(cells) - 1):
                    target.setdefault(label, {})[f"{row}{c}"] = cells[c].strip()
    return layout, results, fit_rows


def gen5_linear(cid: str = "gen5-abs-stdcurve-linear", name: str = "gen5-linear", note: str = "") -> list[Case]:
    layout, results, fits = gen5_text(cid)
    ids, cd = layout["Well ID"], layout["Conc/Dil"]
    raw = {w: float(v) for w, v in results["630nmAbsRead:630"].items()}
    blank = np.mean([raw[w] for w, i in ids.items() if i == "BLK"])
    std = sorted((w for w, i in ids.items() if i.startswith("STD")), key=row_col)
    x = np.array([float(cd[w]) for w in std])
    y = np.array([raw[w] - blank for w in std])
    slope, icpt = np.polyfit(x, y, 1)
    r2 = 1 - np.sum((y - (slope * x + icpt)) ** 2) / np.sum((y - y.mean()) ** 2)
    case = Case(name, {"corpus": cid}, {"analysis": "curve", "model": "linear"},
                "Gen5 BCA-style linear standard curve from the export's own layout; vs numpy and Gen5's results" + note)
    src = f"numpy {np.__version__} polyfit"
    case.add("curve.fit.parameters[name=slope].value", slope, src, rel=1e-9)
    case.add("curve.fit.parameters[name=intercept].value", icpt, src, rel=1e-7, abs_=1e-12)
    case.add("curve.fit.r_squared", r2, src, rel=1e-9)
    head, vals = fits[0]
    a, b = float(vals[head.index("A")]), float(vals[head.index("B")])
    src_v = "Gen5 3.12 StdCurve fit (vendor-computed, 3 significant digits in the file)"
    case.add("curve.fit.parameters[name=slope].value", a, src_v, rel=3e-3)
    case.add("curve.fit.parameters[name=intercept].value", b, src_v, abs_=2e-4)
    src_c = "Gen5 3.12 [Concentration] (vendor-computed from unrounded ODs; the export rounds ODs to 0.001)"
    for w, v in sorted(results["[Concentration]"].items(), key=lambda kv: row_col(kv[0])):
        if re.fullmatch(r"-?\d+(\.\d+)?", v):
            # tolerance: 0.0006 OD of rounding in the exported ODs, divided by the slope
            case.add(f"wells[well={w}].back_calculated", float(v), src_c, rel=5e-3, abs_=0.0006 / a)
    return [case]


def gen5_meanv_4pl(cid: str = "gen5-abs-kinetic-meanv-4pl", name: str = "gen5-meanv-4pl", note: str = "") -> list[Case]:
    layout, results, fits = gen5_text(cid)
    ids, cd = layout["Well ID"], layout["Conc/Dil"]
    mv = {w: float(v) for w, v in results["Mean V [420]"].items()}
    std = sorted((w for w, i in ids.items() if i.startswith("STD")), key=row_col)
    x = np.array([float(cd[w]) for w in std])
    y = np.array([mv[w] for w in std])
    case = Case(name, {"corpus": cid}, {"analysis": "curve", "read": "Mean V [420]"},
                "Gen5 4PL standard curve on the Mean V results (a vendor-calculated read) vs SciPy, drc and Gen5's curve fit" + note)
    p, cov, sse, r2 = fit(four_pl, x, y, [1.0, 3.0, 15.0, 35.0])
    if p[1] < 0:
        p = np.array([p[3], -p[1], p[2], p[0]])
    se = np.sqrt(np.diag(cov))
    for i, n in enumerate("abcd"):
        case.add(f"curve.fit.parameters[name={n}].value", p[i], SCIPY, rel=1e-5)
        case.add(f"curve.fit.parameters[name={n}].se", se[i], SCIPY + " (pcov)", rel=1e-3)
    case.add("curve.fit.r_squared", r2, SCIPY, rel=1e-7)
    head, vals = fits[0]
    src_v = "Gen5 3.12 Curve Fitting Results (vendor-computed, 3 significant digits in the file)"
    for n, col in zip("abcd", "ABCD"):
        case.add(f"curve.fit.parameters[name={n}].value", float(vals[head.index(col)]), src_v, rel=5e-3)
    case.add("curve.fit.r_squared", float(vals[head.index("R2")]), src_v, abs_=6e-4)
    src_c = "Gen5 3.12 [Concentration] (vendor-computed)"
    for w, v in sorted(results["[Concentration]"].items(), key=lambda kv: row_col(kv[0])):
        if re.fullmatch(r"-?\d+(\.\d+)?", v):
            # Gen5 fitted unrounded Mean V values; its parameters are printed to 3 digits. Allow
            # 0.3 % plus the concentration change of 0.2 % of the signal span near the asymptotes.
            x0 = float(v)
            slope = abs((four_pl(x0 * 1.001 + 1e-9, *p) - four_pl(x0 * 0.999, *p)) / (x0 * 0.002 + 1e-9))
            case.add(f"wells[well={w}].back_calculated", x0, src_c, rel=3e-3, abs_=max(2e-3, 0.002 * abs(p[3] - p[0]) / max(slope, 1e-12)))
        elif v.startswith("<"):
            case.add(f"wells[well={w}].flag", ["below_range", "below_curve"], src_c + " (" + v + ")")
        elif v.startswith(">"):
            case.add(f"wells[well={w}].flag", ["above_range", "above_curve"], src_c + " (" + v + ")")
    for w in std[:4]:
        b = inv4(mv[w], *p)
        if b is not None:
            case.add(f"wells[well={w}].back_calculated", b, SCIPY + " inverse 4PL", rel=1e-5)
    if R:
        m = drc(x, y)
        co = m["coef"]
        ours = {"a": co["d"], "b": co["b"], "c": co["e"], "d": co["c"]} if co["b"] > 0 else {"a": co["c"], "b": -co["b"], "c": co["e"], "d": co["d"]}
        for n in "abcd":
            case.add(f"curve.fit.parameters[name={n}].value", ours[n], drc_src() + " LL.4", rel=1e-4)
            # Hessian-based SEs (see the dose-response case): a few percent from Gauss–Newton ones
            case.add(f"curve.fit.parameters[name={n}].se", m["se"][{"a": "d", "d": "c", "c": "e", "b": "b"}[n] if co["b"] > 0 else {"a": "c", "d": "d", "c": "e", "b": "b"}[n]], drc_src() + " vcov, Hessian-based SE", rel=0.1)
        nl = nls4(x, y, p)
        for n in "abcd":
            case.add(f"curve.fit.parameters[name={n}].value", nl["coef"][{"a": "aa", "b": "bb", "c": "cc", "d": "dd"}[n]], nls_src(), rel=1e-6)
            case.add(f"curve.fit.parameters[name={n}].se", nl["se"][{"a": "aa", "b": "bb", "c": "cc", "d": "dd"}[n]], nls_src() + " vcov (Gauss–Newton)", rel=1e-4)
    return [case]


HEADERLESS = " (the export without its file header: the layout embedded after it must still be found)"


def gen5_linear_headerless() -> list[Case]:
    return gen5_linear("synthetic-gen5-headerless-stdcurve-linear", "gen5-linear-headerless", HEADERLESS)


def gen5_meanv_4pl_headerless() -> list[Case]:
    return gen5_meanv_4pl("synthetic-gen5-headerless-kinetic-meanv-4pl", "gen5-meanv-4pl-headerless", HEADERLESS)


def gen5_lum_normalised() -> list[Case]:
    """Gen5's own normalisation of a screening plate (NormLum: % of the POSCON-NEGCON window,
    vendor-computed and stored under each well), from the layout Gen5 embeds (POSCON and NEGCON
    Well IDs are the positive and negative controls)."""
    layout, results, _ = gen5_text("gen5-lum-endpoint")
    ids = layout["Well ID"]
    norm = {w: float(v) for w, v in results["NormLum"].items()}
    case = Case(
        "gen5-lum-normalised",
        {"corpus": "gen5-lum-endpoint"},
        {"analysis": "wells", "normalize": "controls"},
        "Gen5 screening plate: POSCON/NEGCON Well IDs read as positive/negative controls; percent effect vs Gen5's NormLum",
    )
    src_l = "Gen5 3.x embedded layout (Well ID POSCON / NEGCON)"
    case.add("layout.roles.positive", sum(1 for v in ids.values() if v == "POSCON"), src_l)
    case.add("layout.roles.negative", sum(1 for v in ids.values() if v == "NEGCON"), src_l)
    src = "Gen5 3.x NormLum (vendor-computed, 3 decimals in the file)"
    for w, v in sorted(norm.items(), key=lambda kv: row_col(kv[0])):
        # Gen5 prints 3 decimals; its control means differ from the printed integer LUM means
        # by < 0.01 % of the window
        case.add(f"wells[well={w}].percent_effect", v, src, abs_=0.012)
    return [case]


CASES = [
    synth_dose_response,
    synth_elisa,
    synth_kinetics,
    synth_growth,
    skanit,
    gen5_linear,
    gen5_meanv_4pl,
    tecan_growth,
    gen5_linear_headerless,
    gen5_meanv_4pl_headerless,
    gen5_lum_normalised,
]


def render(doc: dict) -> str:
    return json.dumps(doc, indent=1, ensure_ascii=False) + "\n"


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--check", action="store_true", help="exit 1 if a committed case differs")
    ap.add_argument("--only", nargs="*", default=None, help="only these case functions (e.g. gen5_lum_normalised)")
    args = ap.parse_args()
    cases = [m for m in CASES if not args.only or m.__name__ in args.only]
    if not R and not args.only:
        print("warning: no Rscript with drc and growthcurver (OPENREADOUT_RSCRIPT): drc/growthcurver checks are left out", file=sys.stderr)
        if args.check:
            return 1
    bad = 0
    OUT.mkdir(parents=True, exist_ok=True)
    for make in cases:
        for case in make():
            path = OUT / f"{case.name}.json"
            text = render(case.doc())
            if args.check:
                if not path.exists() or path.read_text(encoding="utf-8") != text:
                    print(f"DIFFERS {path.relative_to(ROOT)}")
                    bad += 1
                continue
            path.write_text(text, encoding="utf-8")
            print(f"wrote {path.relative_to(ROOT)}: {len(case.checks)} checks")
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
